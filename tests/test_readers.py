import json
from pathlib import Path

import openpyxl
import pytest
from templify.models import Data
from templify.readers import ExtensionNotSupportedError, ReaderRegistry, readers


class TestRegistry:
    def test_builtin_extensions(self) -> None:
        assert "xlsx" in readers
        assert "json" in readers

    def test_extension_is_case_insensitive(self) -> None:
        assert "XLSX" in readers
        assert readers["JSON"] is readers["json"]

    def test_unsupported_extension_raises(self) -> None:
        assert "csv" not in readers
        with pytest.raises(ExtensionNotSupportedError, match="'csv'"):
            readers["csv"]

    def test_add_reader(self) -> None:
        registry = ReaderRegistry()
        registry.add_reader(lambda _: [{"a": 1}], "fake")
        assert registry["fake"](Path("x.fake")) == [{"a": 1}]

    def test_reader_decorator(self) -> None:
        registry = ReaderRegistry()

        @registry.reader("fake")
        def read_fake(_: Path) -> Data:
            return []

        assert registry["fake"](Path("x.fake")) == []


class TestReadExcel:
    def test_reads_rows_as_dicts(self, xlsx_file: Path, people: Data) -> None:
        assert readers["xlsx"](xlsx_file) == people

    def test_header_only_returns_empty(self, tmp_path: Path) -> None:
        wb = openpyxl.Workbook()
        wb.active.append(["name", "age"])  # type: ignore[union-attr]
        path = tmp_path / "empty.xlsx"
        wb.save(path)
        assert readers["xlsx"](path) == []

    def test_skips_blank_rows(self, tmp_path: Path) -> None:
        wb = openpyxl.Workbook()
        sheet = wb.active
        assert sheet is not None
        sheet.append(["name"])
        sheet.append(["Rama"])
        sheet.append([None])
        sheet.append(["Rawaa"])
        path = tmp_path / "blank.xlsx"
        wb.save(path)
        assert readers["xlsx"](path) == [{"name": "Rama"}, {"name": "Rawaa"}]

    def test_reads_cached_formula_values_not_formulas(self, tmp_path: Path) -> None:
        wb = openpyxl.Workbook()
        sheet = wb.active
        assert sheet is not None
        sheet.append(["name", "total"])
        sheet.append(["Rama", "=1+1"])
        path = tmp_path / "formula.xlsx"
        wb.save(path)
        # openpyxl never computes formulas, so a never-opened file has no cached value
        assert readers["xlsx"](path) == [{"name": "Rama", "total": None}]

    def test_workbook_without_active_sheet_raises(
        self, xlsx_file: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        class NoSheetWorkbook:
            active = None

        monkeypatch.setattr(openpyxl, "load_workbook", lambda *_, **__: NoSheetWorkbook())
        with pytest.raises(ValueError, match="No sheet found"):
            readers["xlsx"](xlsx_file)


class TestReadJson:
    def test_reads_list_of_objects(self, json_file: Path, people: Data) -> None:
        assert readers["json"](json_file) == people

    def test_empty_list(self, tmp_path: Path) -> None:
        path = tmp_path / "empty.json"
        path.write_text("[]")
        assert readers["json"](path) == []

    def test_unicode(self, tmp_path: Path) -> None:
        path = tmp_path / "ar.json"
        path.write_text(json.dumps([{"name": "عبدالله"}], ensure_ascii=False), encoding="utf-8")
        assert readers["json"](path) == [{"name": "عبدالله"}]

    @pytest.mark.parametrize("content", ['{"name": "Rama"}', "[1, 2]", '"text"', "[[]]"])
    def test_not_a_list_of_objects_raises(self, tmp_path: Path, content: str) -> None:
        path = tmp_path / "bad.json"
        path.write_text(content)
        with pytest.raises(ValueError, match="list of objects"):
            readers["json"](path)

    def test_malformed_json_raises(self, tmp_path: Path) -> None:
        path = tmp_path / "bad.json"
        path.write_text("[{")
        with pytest.raises(json.JSONDecodeError):
            readers["json"](path)
