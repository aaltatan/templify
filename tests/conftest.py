import json
from collections.abc import Callable
from pathlib import Path

import docx
import openpyxl
import pytest
from templify.models import Data

type MakeFile = Callable[[str, str], Path]


@pytest.fixture
def people() -> Data:
    return [
        {"name": "Abdullah", "age": 32},
        {"name": "Rawaa", "age": 33},
        {"name": "Rama", "age": 25},
    ]


@pytest.fixture
def make_text_file(tmp_path: Path) -> MakeFile:
    def make(filename: str, content: str) -> Path:
        path = tmp_path / filename
        path.write_text(content, encoding="utf-8")
        return path

    return make


@pytest.fixture
def make_docx(tmp_path: Path) -> MakeFile:
    """Create a docx template, each line of `content` becomes a paragraph."""

    def make(filename: str, content: str) -> Path:
        document = docx.Document()
        for line in content.splitlines():
            document.add_paragraph(line)
        path = tmp_path / filename
        document.save(str(path))
        return path

    return make


@pytest.fixture
def xlsx_file(tmp_path: Path, people: Data) -> Path:
    wb = openpyxl.Workbook()
    sheet = wb.active
    assert sheet is not None
    sheet.append(list(people[0]))
    for person in people:
        sheet.append(list(person.values()))
    path = tmp_path / "people.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def json_file(tmp_path: Path, people: Data) -> Path:
    path = tmp_path / "people.json"
    path.write_text(json.dumps(people), encoding="utf-8")
    return path


@pytest.fixture
def output_dir(tmp_path: Path) -> Path:
    return tmp_path / "output"


def read_docx(path: Path) -> str:
    return "\n".join(p.text for p in docx.Document(str(path)).paragraphs)
