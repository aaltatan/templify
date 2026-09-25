"""Data readers: turn a data file into a list of records, picked by file extension."""

import json
from collections.abc import Callable
from functools import wraps
from pathlib import Path
from typing import TypeGuard

import openpyxl

from .models import Data

type ReadFn = Callable[[Path], Data]


class ExtensionNotSupportedError(Exception):
    """Raised when no reader is registered for a data file extension."""

    def __init__(self, extension: str) -> None:
        message = f"Extension '{extension}' is not supported for being a data file"
        super().__init__(message)


class ReaderRegistry:
    """Maps data file extensions (case-insensitive, without the dot) to read functions.

    Example:
    ```python
    readers = ReaderRegistry()


    @readers.reader("csv")
    def read_csv(filepath: Path) -> Data: ...


    data = readers["csv"](Path("people.csv"))
    ```

    """

    def __init__(self) -> None:
        self._read_fns: dict[str, ReadFn] = {}

    def __contains__(self, extension: str) -> bool:
        """Whether a reader is registered for `extension`."""
        return extension.lower() in self._read_fns

    def __getitem__(self, extension: str) -> ReadFn:
        """Return the reader registered for `extension`.

        Raises:
            ExtensionNotSupportedError: No reader is registered for `extension`.

        """
        if extension not in self:
            raise ExtensionNotSupportedError(extension)
        return self._read_fns[extension.lower()]

    def reader(self, extension: str) -> Callable[[ReadFn], ReadFn]:
        """Register the decorated function as the reader for `extension`."""

        def decorator(fn: ReadFn) -> ReadFn:
            @wraps(fn)
            def wrapper(filepath: Path) -> Data:
                return fn(filepath)

            self._add_reader(wrapper, extension)

            return wrapper

        return decorator

    def add_reader(self, fn: ReadFn, extension: str) -> None:
        """Register `fn` as the reader for `extension`, replacing any existing one."""
        self._add_reader(fn, extension)

    def _add_reader(self, fn: ReadFn, extension: str) -> None:
        self._read_fns[extension.lower()] = fn


readers = ReaderRegistry()
"""The default registry, used by the CLI `--data` option."""


@readers.reader("xlsx")
def read_excel(filepath: Path) -> Data:
    """Read the active sheet of an Excel workbook.

    The first row holds the column names, each following row is one record.
    Fully blank rows are skipped, and formula cells give their last calculated value.

    Raises:
        ValueError: The workbook has no active sheet.

    """
    wb = openpyxl.load_workbook(filepath, data_only=True)
    sheet = wb.active

    if sheet is None:
        message = "No sheet found"
        raise ValueError(message)

    headers = []

    for row in sheet.iter_rows(min_row=1, max_row=1):
        for cell in row:
            headers.append(cell.value)  # noqa: PERF401

    data = []

    for row in sheet.iter_rows(min_row=2):
        if all(cell.value is None for cell in row):
            continue

        row_dict = {}

        for idx, cell in enumerate(row):
            row_dict[headers[idx]] = cell.value

        data.append(row_dict)

    return data


@readers.reader("json")
def read_json(filepath: Path) -> Data:
    """Read a UTF-8 JSON file holding a list of objects.

    Raises:
        ValueError: The file is not valid JSON, or not a list of objects.

    """
    with filepath.open(encoding="utf-8") as f:
        data = json.load(f)

    if _is_data(data):
        return data

    message = "JSON data must be a list of objects"
    raise ValueError(message)


def _is_data(d: object) -> TypeGuard[Data]:
    return (
        isinstance(d, list)
        and all(isinstance(item, dict) for item in d)
        and all(all(isinstance(key, str) for key in item) for item in d)
    )
